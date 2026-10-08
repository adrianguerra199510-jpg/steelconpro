# -*- coding: utf-8 -*-
"""Reportes: imagen, hoja de calculo y memoria de calculo en Word."""
from __future__ import annotations
from pathlib import Path
import datetime
import math

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from .model import Project
from .solver import Results
from . import draw, geometry as G
from .units import U, UnitSet, float_to_frac


def _units(prj) -> UnitSet:
    return UnitSet(getattr(prj, "u_len", "in"), getattr(prj, "u_force", "kip"),
                   getattr(prj, "u_stress", "ksi"), getattr(prj, "u_moment", None))


def ck_vals(us: UnitSet, ch):
    """Demanda, capacidad y etiqueta de unidad de una verificacion, ya
    convertidas al sistema que eligio el usuario."""
    k = ch.kind
    if k == "-":
        return ch.demand, ch.capacity, ch.unit
    return us.out(k, ch.demand), us.out(k, ch.capacity), us.label(k)


# =================================================================== imagenes
def input_rows(prj: Project, us: UnitSet) -> list[tuple[str, str]]:
    """Tabla de datos de entrada, convertida a las unidades elegidas."""
    p, b, c, L = prj.plate, prj.bolts, prj.conc, prj.eloads
    q = us.q
    geo = (f"Ø{q('L', p.Dp)}" if p.shape == "Circular"
           else f"{q('L', p.N)} × {q('L', p.B)}")
    return [
        ("Perfil", f"{prj.section.label} ({prj.section.steel}), "
                   f"rotacion {prj.section.rotation:g}°"),
        ("Placa base", f"{geo} × {q('L', p.tp)}, {p.steel} "
                       f"(Fy = {q('S', p.mat().Fy)})"),
        ("Mortero", q("L", p.grout)),
        ("Anclajes", f"{b.n_total} Ø{b.size} in ({q('L', b.geom().db)}), {b.steel}, "
                     f"{b.atype}, hef = {q('L', b.hef)}"),
        ("Agujeros", f"{q('L', b.geom().dh)} — {b.hole_rule}"),
        ("Disposicion", f"{b.pattern} — {b.n_major} en eje mayor, "
                        f"{b.n_minor} en eje menor, ex = {q('L', b.ex)}, "
                        f"ey = {q('L', b.ey)}"),
        ("Concreto", f"f'c = {q('S', c.fc)}, pedestal {q('L', c.N2)} × {q('L', c.B2)}, "
                     f"ha = {q('L', c.ha)}, "
                     f"{'fisurado' if c.cracked else 'no fisurado'}, "
                     f"condicion {'A' if c.cond_A_eff else 'B'}"),
        ("Llave de corte", (f"{q('L', prj.lug.W)} × {q('L', prj.lug.H)} × "
                            f"{q('L', prj.lug.t)}, {prj.lug.direction}"
                            if prj.lug.enabled else "No")),
        ("Rigidizadores", (f"{prj.stiff.count} pletinas {q('L', prj.stiff.t)} × "
                           f"{q('L', prj.stiff.h)}, proyeccion {q('L', prj.stiff.L)}, "
                           f"{prj.stiff.shape}, {prj.stiff.position}"
                           if prj.stiff.enabled else "No")),
        ("Soldadura ala", f"{prj.welds.flange.wtype} {q('L', prj.welds.flange.size)}, "
                          f"{prj.welds.flange.electrode}"),
        ("Soldadura alma", f"{prj.welds.web.wtype} {q('L', prj.welds.web.size)}, "
                           f"{prj.welds.web.electrode}"),
        ("Soldadura perimetral", f"{prj.welds.perimeter.wtype} "
                                 f"{q('L', prj.welds.perimeter.size)}, "
                                 f"{prj.welds.perimeter.electrode}"),
        ("Cargas (LRFD)" + (" en ejes de la placa" if prj.loads.tilted else ""),
         f"Pu = {q('F', L.Pu)}, Mux = {q('M', L.Mux)}, "
         f"Muy = {q('M', L.Muy)}, Vux = {q('F', L.Vux)}, "
         f"Vuy = {q('F', L.Vuy)}"),
        ("Columna inclinada", (f"giro X = {prj.loads.tilt_x:g}°, giro Y = {prj.loads.tilt_y:g}°; "
                               f"cargas en el eje: Pu = {q('F', prj.loads.Pu)}, "
                               f"Vux = {q('F', prj.loads.Vux)}, Vuy = {q('F', prj.loads.Vuy)}, "
                               f"Mux = {q('M', prj.loads.Mux)}, Muy = {q('M', prj.loads.Muy)}")
                              if prj.loads.tilted else "no (perpendicular a la placa)"),
    ]


def _fem_of(res):
    return getattr(res, "fem", None)


def bolt_rows(prj: Project, res: Results, us: UnitSet):
    """(encabezados, filas) de la tabla de traccion por perno del modelo 3D."""
    fem = _fem_of(res)
    post = getattr(fem, "post", None)
    if post is None or not getattr(post, "bolts", None):
        return None, []
    from .fem_checks import bolt_phiRnt
    phi = bolt_phiRnt(prj)
    hdr = ["Perno", f"x ({us.L})", f"y ({us.L})", f"T ({us.F})", "D/C"]
    filas = []
    for (k, x, y, T) in sorted(post.bolts, key=lambda b_: -b_[3]):
        filas.append([f"P{k}", us.fmt("L", x), us.fmt("L", y), us.fmt("F", T),
                      f"{T / phi:.3f}" if phi else "—"])
    return hdr, filas


def fem_section(prj: Project, res: Results, us: UnitSet):
    """Contenido de las secciones del modelo solido 3D.  -> None si no hay un 3D vigente."""
    fem = _fem_of(res)
    if fem is None or fem.post is None:
        return None
    from .weld3d import summary_rows
    post, rep = fem.post, (fem.rep or {})
    conn = getattr(post, "weld_model", "") == "conectores"
    intro = ("Modelo solido de tetraedros cuadraticos (Gmsh + CalculiX): placa con los agujeros "
             "taladrados, perfil, rigidizadores y llave; concreto como resortes de Winkler solo a "
             "compresion y pernos como resortes solo a traccion (paso no lineal). ")
    if conn:
        intro += ("SOLDADURA: el perfil y la placa se modelan como cuerpos separados. La compresion pasa por "
                  "contacto (resortes solo-compresion en toda la huella) y cada linea de cordon (cara exterior "
                  "o interior de cada pared, y la base de los rigidizadores) es un conector de traccion y "
                  "cortante con la rigidez de la garganta; su fuerza se lee directamente del resorte, y un "
                  "lado sin cordon no transmite carga. La fuerza por unidad de longitud se suaviza en una "
                  "ventana movil de 4 veces el cateto (redistribucion por ductilidad del cordon) y se "
                  "compara, por linea, con φ·0.60·FEXX·garganta·kd (AISC J2.4) y, con la suma de las lineas "
                  "de la pared, con la rotura del metal base. D/C pico = maximo de la curva suavizada; "
                  "D/C media = fuerza de la linea repartida en su longitud.")
    else:
        intro += ("SOLDADURA: union monolitica (equivale a CJP); la fuerza del cordon se obtiene integrando en "
                  "el espesor de cada pared los esfuerzos del perfil justo por encima del pie del cordon y se "
                  "verifica con el metodo vectorial de AISC J2.4. D/C pico = punto mas cargado; D/C media = "
                  "fuerza de la pared repartida en su longitud.")
    if conn:
        intro += (f" Compresion por contacto perfil-placa = {us.q('F', post.F_bear)}; traccion tomada por los "
                  f"cordones = {us.q('F', post.Fz_weld)}.")
    weld_rows, _ = summary_rows(prj, post)
    b_hdr, b_rows = bolt_rows(prj, res, us)
    plan = rep.get("plan") or {}
    figs = [(plan.get("top"), "Placa aislada, cara superior: von Mises promediado (etiqueta: esfuerzo maximo)"),
            (plan.get("bot"), "Placa aislada, cara inferior (apoyada en el mortero): von Mises promediado"),
            (plan.get("press"), "Presion de contacto sobre el concreto y traccion en cada perno"),
            (plan.get("uz"), "Desplazamiento vertical de la cara superior de la placa")]
    figs = [(a, b) for a, b in figs if a]
    persp = [(rep.get("vm"), "Vista 3D: esfuerzo de von Mises (etiqueta: esfuerzo maximo)"),
             (rep.get("u"), "Vista 3D: desplazamiento |U| sobre la geometria deformada")]
    persp = [(a, b) for a, b in persp if a]
    return dict(intro=intro, msg=f"Equilibrio: {post.msg}.", weld_rows=weld_rows, b_hdr=b_hdr, b_rows=b_rows,
                figs=figs, persp=persp, lines=_rep3d_lines(prj, us, rep) if rep else [])


def _combo_name(prj: Project) -> str:
    try:
        return prj.combo_list()[prj.combo_idx].name
    except Exception:
        return ""


def render_3d_png(prj: Project, path: str, size=(6.8, 5.2), dpi=170) -> str | None:
    """Vista 3D de la geometria con las flechas de las cargas de la combinacion para la memoria de calculo."""
    try:
        from . import view3d
        fig = plt.figure(figsize=size, dpi=dpi)
        ax = fig.add_axes([0.0, 0.0, 1.0, 0.94], projection="3d")
        view3d.plot_geometry(ax, prj, loads=True, title=False)
        ax.view_init(elev=24, azim=-58)
        view3d.update_order(ax)
        view3d.fit_to_axes(ax)
        fig.savefig(path)
        plt.close(fig)
        return path
    except Exception:
        plt.close("all")
        return None


def _rep3d_lines(prj: Project, us, rep: dict) -> list[str]:
    """Texto del resumen de resultados del modelo solido 3D."""
    out = []
    if rep.get("lc"):
        out.append(f"Tamano de elemento: {us.q('L', rep['lc'])} ("
                   + ("calculado automaticamente: el mayor entre 1.2 veces el radio de promedio y raiz(area de "
                      "placa/400); en el estudio de convergencia el von Mises promediado varia ±2 % y la "
                      "traccion en pernos < 0.5 %)." if rep.get("fast") else "definido por el usuario).")
                   )
    out.append(f"Modelo de {rep['n_nodes']:,} nodos y {rep['n_elems']:,} tetraedros.")
    va = rep.get("vm_avg")
    if va:
        out.append(f"ESFUERZO MAXIMO (von Mises promediado en la placa, r = "
                   f"{us.q('L', va['radius'])}) = {us.q('S', va['vm'])}, en "
                   f"({us.fmt('L', va['x'])}, {us.fmt('L', va['y'])}) (marcado con ★ en la "
                   f"figura). Pico puntual = {us.q('S', rep['vmmax'])}, que crece al refinar "
                   f"la malla y no debe usarse para verificar.")
    else:
        out.append(f"ESFUERZO MAXIMO (von Mises, pico puntual) = {us.q('S', rep['vmmax'])}"
                   + (f", en el nodo {rep['vm_node']} (marcado con ★ en la figura)"
                      if rep.get("vm_node") else "") + ".")
    out.append(f"Desplazamiento maximo |U| = {us.q('L', rep['umax'])}"
               + (f"; deformada amplificada ×{rep['scale']:g}." if rep.get("scale") else "."))
    out.append("El promedio se calcula sobre el tensor de esfuerzos, ponderado por el area de "
               "cada nodo, dentro de un circulo de radio fijo sobre la misma cara de la placa "
               "(la cara superior excluye lo cubierto por el perfil): converge al refinar la "
               "malla. Los picos puntuales en aristas vivas (borde de agujero, encuentro "
               "perfil-placa) son singularidades de malla.")
    return out


_FEM_PNGS = ("planta_vm_sup.png", "planta_vm_inf.png", "planta_presion.png", "planta_uz.png")


def annex_figs(figs):
    """Dibujos del anexo grafico: los mapas del 3D ya estan en su seccion, no se repiten."""
    return [fp for fp in (figs or []) if Path(fp).name not in _FEM_PNGS]


def save_figures(prj: Project, res: Results, folder: str) -> list[str]:
    f = Path(folder)
    f.mkdir(parents=True, exist_ok=True)
    paths = []

    fig, ax = plt.subplots(figsize=(7.0, 7.0), dpi=150)
    draw.plan_view(ax, prj)
    fig.tight_layout()
    p1 = f / "planta.png"; fig.savefig(p1); plt.close(fig); paths.append(str(p1))

    fig, ax = plt.subplots(figsize=(7.0, 5.0), dpi=150)
    draw.elevation_view(ax, prj)
    fig.tight_layout()
    p2 = f / "elevacion.png"; fig.savefig(p2); plt.close(fig); paths.append(str(p2))

    fs = fem_section(prj, res, _units(prj))
    if fs:
        # placa aislada en planta con los resultados del modelo solido 3D
        for i, (fp, cap) in enumerate(fs["figs"]):
            try:
                import shutil
                dst = f / Path(fp).name
                if Path(fp).resolve() != dst.resolve():
                    shutil.copyfile(fp, dst)
                paths.append(str(dst))
            except Exception:
                pass
    return paths


# ==================================================================== XLSX
def export_xlsx(prj: Project, res: Results, path: str,
                figs: list[str] | None = None, detail: bool = True) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage

    u = U(prj.metric)
    F = "Arial"
    B = Font(name=F, size=10, bold=True)
    N = Font(name=F, size=10)
    S = Font(name=F, size=8, italic=True, color="595959")
    HD = Font(name=F, size=10, bold=True, color="FFFFFF")
    FH = PatternFill("solid", fgColor="2E75B6")
    OKF = PatternFill("solid", fgColor="C6EFCE")
    NOF = PatternFill("solid", fgColor="FFC7CE")
    thin = Side(style="thin", color="BFBFBF")
    BX = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "RESUMEN"
    for col, w in zip("ABCDEFG", (3, 54, 14, 14, 10, 11, 74)):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    def band(r, t):
        ws.merge_cells(f"B{r}:G{r}")
        ws[f"B{r}"] = t
        for c in "BCDEFG":
            ws[f"{c}{r}"].fill = FH
            ws[f"{c}{r}"].font = HD
        ws[f"B{r}"].alignment = Alignment(horizontal="left", indent=1)

    ws["B1"] = "PLACA BASE — RESUMEN DE VERIFICACIONES"
    ws["B1"].font = Font(name=F, size=14, bold=True)
    ws["B2"] = (f"{prj.name}   |   Elemento {prj.element}   |   {prj.author}   |   "
                f"{prj.date or datetime.date.today().isoformat()}")
    ws["B2"].font = S

    r = 4
    band(r, "1.  DATOS PRINCIPALES"); r += 1
    p, b, c, L = prj.plate, prj.bolts, prj.conc, prj.eloads
    geo = (f"Ø{u.l(p.Dp):.4g}" if p.shape == "Circular"
           else f"{u.l(p.N):.4g} × {u.l(p.B):.4g}")
    datos = [
        ("Perfil", f"{prj.section.label}  ({prj.section.steel}, rotacion {prj.section.rotation:g}°)"),
        ("Placa", f"{geo} × {u.l(p.tp):.4g} {u.L}  —  {p.steel}"),
        ("Mortero de nivelacion", f"{u.l(p.grout):.4g} {u.L}"),
        ("Pernos", f"{b.n_total} × Ø{b.size}\"  {b.steel}  —  {b.atype}"),
        ("Disposicion", f"{b.pattern};  eje mayor {b.n_major} / eje menor {b.n_minor};  "
                        f"hef = {u.l(b.hef):.4g} {u.L}"),
        ("Concreto", f"f'c = {u.s(c.fc):.4g} {u.S};  pedestal {u.l(c.N2):.4g} × {u.l(c.B2):.4g} {u.L};  "
                     f"{'fisurado' if c.cracked else 'no fisurado'}"),
        ("Llave de corte", (f"{u.l(prj.lug.W):.4g} × {u.l(prj.lug.H):.4g} × {u.l(prj.lug.t):.4g} {u.L}"
                            if prj.lug.enabled else "no")),
        ("Rigidizadores", (f"{prj.stiff.count} pletinas {u.l(prj.stiff.t):.4g} {u.L}, "
                           f"L = {u.l(prj.stiff.L):.4g} {u.L}" if prj.stiff.enabled else "no")),
        ("Soldadura", f"ala: {prj.welds.flange.wtype} {float_to_frac(prj.welds.flange.size)}\" · "
                      f"alma: {prj.welds.web.wtype} {float_to_frac(prj.welds.web.size)}\" · "
                      f"perim.: {prj.welds.perimeter.wtype} {float_to_frac(prj.welds.perimeter.size)}\""),
        ("Columna inclinada", (f"giro X = {prj.loads.tilt_x:g}°, giro Y = {prj.loads.tilt_y:g}° "
                               f"(cargas ingresadas en el eje de la columna)"
                               if prj.loads.tilted else "no")),
        ("Cargas factorizadas" + (" (ejes de la placa)" if prj.loads.tilted else ""),
         f"Pu = {u.f(L.Pu):.4g} {u.F};  Mux = {u.m(L.Mux):.4g} {u.M};  "
                                f"Muy = {u.m(L.Muy):.4g} {u.M};  Vu = {u.f(L.Vu):.4g} {u.F}"),
    ]
    for k, v in datos:
        ws[f"B{r}"] = k; ws[f"B{r}"].font = N
        ws.merge_cells(f"C{r}:G{r}")
        ws[f"C{r}"] = v; ws[f"C{r}"].font = N
        r += 1

    r += 1
    band(r, "2.  ESTADO DE APLASTAMIENTO"); r += 1
    br = res.br
    for k, v in [("Caso", br.case),
                 ("Excentricidad e = Mu/Pu", f"{u.l(br.e):.4g} {u.L}" if br.e != float('inf') else "∞"),
                 ("e critica", f"{u.l(br.ecrit):.4g} {u.L}"),
                 ("Longitud de aplastamiento Y", f"{u.l(br.Y):.4g} {u.L}"),
                 ("Presion de contacto fp", f"{u.s(br.fp):.4g} / {u.s(br.fp_max):.4g} {u.S}"),
                 ("Traccion total en pernos Tu", f"{u.f(br.Tu):.4g} {u.F}  en {br.n_t} pernos"),
                 ("Brazo del grupo traccionado f", f"{u.l(br.f_arm):.4g} {u.L}"),
                 ("Espesor requerido", f"{u.l(res.treq):.4g} / {u.l(p.tp):.4g} {u.L}")]:
        ws[f"B{r}"] = k; ws[f"B{r}"].font = N
        ws.merge_cells(f"C{r}:G{r}"); ws[f"C{r}"] = v; ws[f"C{r}"].font = N
        r += 1

    r += 1
    band(r, "3.  VERIFICACIONES"); r += 1
    hdr = ["Verificacion", "Demanda", "Capacidad", "Unid.", "D/C", "Referencia y observaciones"]
    for i, h in enumerate(hdr):
        cell = ws.cell(row=r, column=2 + i, value=h)
        cell.font = B; cell.border = BX
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
    r += 1
    first = r
    usx = _units(prj)
    for ch in res.checks:
        dv, cv, ul = ck_vals(usx, ch)
        ws.cell(row=r, column=2, value=ch.title).font = N
        ws.cell(row=r, column=3, value=round(dv, 4)).font = N
        ws.cell(row=r, column=4, value=round(cv, 4)).font = N
        ws.cell(row=r, column=5, value=ul).font = N
        cr = ws.cell(row=r, column=6, value=("—" if ch.skip else round(ch.ratio, 3)))
        cr.font = B
        cr.fill = OKF if ch.ok else NOF
        cr.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=7, value=(ch.ref + ("  —  " + ch.note if ch.note else ""))).font = S
        for cc in range(2, 8):
            ws.cell(row=r, column=cc).border = BX
        r += 1

    r += 1
    ws[f"B{r}"] = "VEREDICTO"
    ws[f"B{r}"].font = Font(name=F, size=12, bold=True)
    gov = res.governing
    ws[f"C{r}"] = "PENDIENTE — realizar analisis 3D" if res.pending else ("CUMPLE" if res.ok else "NO CUMPLE")
    ws[f"C{r}"].font = Font(name=F, size=12, bold=True)
    ws[f"C{r}"].fill = PatternFill("solid", fgColor="FFEB9C") if res.pending else (OKF if res.ok else NOF)
    ws.merge_cells(f"D{r}:G{r}")
    ws[f"D{r}"] = (f"D/C maximo = {res.max_ratio:.3f}" +
                   (f"  —  gobierna: {gov.title}" if gov else ""))
    ws[f"D{r}"].font = N
    r += 2
    if res.warnings:
        band(r, "4.  AVISOS"); r += 1
        for wmsg in res.warnings:
            ws.merge_cells(f"B{r}:G{r}")
            ws[f"B{r}"] = wmsg
            ws[f"B{r}"].font = Font(name=F, size=9, bold=wmsg.startswith("**"),
                                    color="9C0006" if wmsg.startswith("**") else "595959")
            r += 1

    # --- hoja de dibujos
    if figs:
        ws2 = wb.create_sheet("DIBUJOS")
        ws2.sheet_view.showGridLines = False
        row = 2
        for fp in figs:
            try:
                img = XLImage(fp)
                img.width = int(img.width * 0.55)
                img.height = int(img.height * 0.55)
                ws2.add_image(img, f"B{row}")
                row += 32
            except Exception:
                pass

    # --- hoja FEM 3D
    fem = getattr(res, "fem", None)
    if fem is not None and fem.post is not None:
        ws3 = wb.create_sheet("FEM 3D")
        ws3.sheet_view.showGridLines = False
        ws3.column_dimensions["B"].width = 46
        ws3.column_dimensions["C"].width = 18
        post = fem.post
        va = fem.vm_avg
        rows = [("Nodos", fem.n_nodes), ("Tetraedros", fem.n_elems),
                (f"Desplazamiento maximo |U| ({usx.L})", usx.out("L", fem.umax)),
                (f"Presion de contacto maxima ({usx.S})", usx.out("S", post.p_max)),
                (f"von Mises puntual maximo ({usx.S}) — no converge", usx.out("S", fem.vmmax)),
                (f"Reaccion del concreto ({usx.F})", usx.out("F", post.R_conc)),
                (f"Traccion total en pernos ({usx.F})", usx.out("F", post.T_bolts))]
        if va:
            rows.insert(4, (f"von Mises PROMEDIADO maximo ({usx.S}), r = {usx.out('L', va['radius']):.3g} {usx.L}",
                            usx.out("S", va["vm"])))
        rr = 2
        ws3["B1"] = "MODELO SOLIDO 3D — RESULTADOS"; ws3["B1"].font = Font(name=F, size=12, bold=True)
        for k, v in rows:
            ws3[f"B{rr}"] = k; ws3[f"B{rr}"].font = N
            ws3[f"C{rr}"] = v; ws3[f"C{rr}"].font = B
            rr += 1
        rr += 1
        ws3[f"B{rr}"] = "TRACCION PERNO POR PERNO"; ws3[f"B{rr}"].font = B
        rr += 1
        hdr, filas = bolt_rows(prj, res, usx)
        if hdr:
            for j, htxt in enumerate(hdr):
                cc = ws3.cell(row=rr, column=2 + j, value=htxt); cc.font = B
                cc.fill = PatternFill("solid", fgColor="DDDDDD")
            rr += 1
            red = PatternFill("solid", fgColor="FFC7CE")
            grn = PatternFill("solid", fgColor="C6EFCE")
            for fila in filas:
                for j, vtxt in enumerate(fila):
                    try:
                        val = float(str(vtxt).replace(",", ""))
                    except ValueError:
                        val = vtxt
                    cc = ws3.cell(row=rr, column=2 + j, value=val)
                    cc.font = N
                    if j == 4 and isinstance(val, float):
                        cc.font = B
                        cc.fill = grn if val <= 1.0 else red
                rr += 1

    if detail and res.rec is not None:
        ws4 = wb.create_sheet("MEMORIA")
        ws4.sheet_view.showGridLines = False
        ws4.column_dimensions["B"].width = 130
        ws4["B1"] = "DESARROLLO DE LAS ECUACIONES"
        ws4["B1"].font = Font(name=F, size=12, bold=True)
        rr = 3
        for kind, txt in res.rec.to_lines():
            c = ws4.cell(row=rr, column=2, value=txt)
            if kind == "sec":
                c.font = HD; c.fill = FH
            elif kind == "chk":
                c.font = B
            elif kind == "txt":
                c.font = S
            else:
                c.font = N
            rr += 1

    wb.save(path)
    return path


# ==================================================================== DOCX
def export_docx(prj: Project, res: Results, path: str,
                figs: list[str] | None = None, detail: bool = True) -> str:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    u = U(prj.metric)
    us = _units(prj)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9)

    try:
        from . import brand
        doc.add_picture(brand.LOGO(), width=Inches(2.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    except Exception:
        pass
    doc.add_heading("MEMORIA DE CALCULO — PLACA BASE", level=0)
    p0 = doc.add_paragraph()
    p0.add_run(f"{prj.name}\n").bold = True
    p0.add_run(f"Elemento: {prj.element}     Calculo: {prj.author}     "
               f"Fecha: {prj.date or datetime.date.today().isoformat()}\n")
    p0.add_run("Normas: AISC 360-22, AISC Design Guide 1 (2ª Ed.), ACI 318-19 Cap. 17.")

    doc.add_heading("1. Datos de entrada", level=1)
    import tempfile
    _tmp3d = render_3d_png(prj, str(Path(tempfile.mkdtemp()) / "vista3d.png"))
    if _tmp3d:
        doc.add_picture(_tmp3d, width=Inches(4.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = doc.add_paragraph(f"Vista 3D de la conexion con las cargas de la combinacion {_combo_name(prj)}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
    p, b, c, L = prj.plate, prj.bolts, prj.conc, prj.eloads
    rows = input_rows(prj, us)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = v

    doc.add_heading("2. Aplastamiento y equilibrio (DG1 §3.3)", level=1)
    br = res.br
    doc.add_paragraph(
        f"A1 = {br.A1:.1f} in²;  A2 = {br.A2:.1f} in²;  √(A2/A1) = {br.sqrt_ratio:.3f};  "
        f"φcPp = {us.q('F', br.phiPp)};  fp,max = {us.q('S', br.fp_max)};  "
        f"qmax = {us.q('LF', br.qmax)}.")
    doc.add_paragraph(
        f"e = Mu/Pu = {br.e:.2f} in;  ecrit = {br.ecrit:.2f} in  →  {br.case}.  "
        f"Y = {us.q('L', br.Y)};  fp = {us.q('S', br.fp)};  Tu = {us.q('F', br.Tu)} repartidos en "
        f"{br.n_t} pernos con brazo f = {br.f_arm:.2f} in.")
    d = res.tdet
    doc.add_paragraph(
        f"Voladizos: m = {d['m_y']:.2f} in, n = {d['m_x']:.2f} in, λ = {d['lam']:.3f}, "
        f"λn' = {d['lam_n']:.2f} in.  Espesor requerido t = {res.treq:.3f} in "
        f"(propuesto {p.tp:g} in).")

    doc.add_heading("3. Verificaciones", level=1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Verificacion", "Demanda", "Capacidad", "Un.", "D/C"]):
        t.rows[0].cells[i].text = h
        for pr in t.rows[0].cells[i].paragraphs:
            for rr in pr.runs:
                rr.bold = True
    us = _units(prj)
    for ch in res.checks:
        dv, cv, ul = ck_vals(us, ch)
        cells = t.add_row().cells
        cells[0].text = ch.title
        cells[1].text = f"{dv:,.3f}"
        cells[2].text = f"{cv:,.3f}"
        cells[3].text = ul
        cells[4].text = "—" if ch.skip else f"{ch.ratio:.3f}"
        if not ch.ok:
            for pr in cells[4].paragraphs:
                for rr in pr.runs:
                    rr.bold = True
                    rr.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)

    pv = doc.add_paragraph()
    pv.add_run("VEREDICTO: ").bold = True
    rr = pv.add_run("PENDIENTE — realizar el analisis 3D" if res.pending else ("CUMPLE" if res.ok else "NO CUMPLE"))
    rr.bold = True
    rr.font.color.rgb = (RGBColor(0x9C, 0x57, 0x00) if res.pending else
                         (RGBColor(0x00, 0x61, 0x00) if res.ok else RGBColor(0x9C, 0x00, 0x06)))
    gov = res.governing
    pv.add_run(f"   D/C maximo = {res.max_ratio:.3f}" +
               (f"   (gobierna: {gov.title})" if gov else ""))

    if len(getattr(res, "combo_rows", []) or []) > 1:
        pc = doc.add_paragraph()
        pc.add_run("Combinaciones de carga (el detalle de esta memoria corresponde a la que gobierna).").bold = True
        tc = doc.add_table(rows=1, cols=4)
        tc.style = "Light Grid Accent 1"
        for i, h in enumerate(["Combinacion", "D/C max", "Gobierna", "Estado"]):
            tc.rows[0].cells[i].text = h
        for cr in res.combo_rows:
            cells = tc.add_row().cells
            cells[0].text = cr["name"]
            cells[1].text = f"{cr['ratio']:.3f}"
            cells[2].text = cr["gov"]
            cells[3].text = "PENDIENTE" if cr["pending"] else ("CUMPLE" if cr["ok"] else "NO CUMPLE")

    if res.warnings:
        doc.add_heading("4. Avisos", level=1)
        for wmsg in res.warnings:
            doc.add_paragraph(wmsg, style="List Bullet")

    _fs = fem_section(prj, res, us)
    if _fs:
        doc.add_heading("6. Modelo solido 3D — pernos y soldadura", level=1)
        doc.add_paragraph(_fs["intro"])
        doc.add_paragraph(_fs["msg"])
        doc.add_paragraph().add_run("Soldadura perfil-placa").bold = True
        rows_w = _fs["weld_rows"]
        tb = doc.add_table(rows=0, cols=len(rows_w[0]))
        tb.style = "Light Grid Accent 1"
        for i, r in enumerate(rows_w):
            cel = tb.add_row().cells
            for j, v in enumerate(r):
                cel[j].text = ""
                run = cel[j].paragraphs[0].add_run(str(v))
                run.font.size = Pt(8)
                run.bold = (i == 0)
        if _fs["b_hdr"]:
            doc.add_paragraph().add_run("Traccion por perno").bold = True
            tb = doc.add_table(rows=1, cols=len(_fs["b_hdr"]))
            tb.style = "Light Grid Accent 1"
            for j, htxt in enumerate(_fs["b_hdr"]):
                tb.rows[0].cells[j].text = ""
                tb.rows[0].cells[j].paragraphs[0].add_run(htxt).bold = True
            for fila in _fs["b_rows"]:
                cel = tb.add_row().cells
                for j, vtxt in enumerate(fila):
                    cel[j].text = ""
                    cel[j].paragraphs[0].add_run(vtxt).font.size = Pt(8)
        doc.add_heading("7. Placa base aislada — resultados del modelo 3D en planta", level=1)
        for i, ln in enumerate(_fs["lines"]):
            run = doc.add_paragraph().add_run(ln)
            if ln.startswith("ESFUERZO MAXIMO"):
                run.bold = True
        for fp, cap in _fs["figs"] + _fs["persp"]:
            try:
                doc.add_picture(fp, width=Inches(5.6))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                c_ = doc.add_paragraph(cap)
                c_.alignment = WD_ALIGN_PARAGRAPH.CENTER
                c_.runs[0].italic = True
            except Exception:
                pass

    if detail and res.rec is not None:
        doc.add_page_break()
        doc.add_heading("Anexo A — Desarrollo de las ecuaciones", level=1)
        doc.add_paragraph(
            "Cada linea muestra el simbolo, la formula, la sustitucion numerica y "
            "el resultado, en las unidades de trabajo del proyecto.")
        for kind, txt in res.rec.to_lines():
            par = doc.add_paragraph()
            run = par.add_run(txt)
            run.font.size = Pt(7.5)
            if kind == "sec":
                run.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            elif kind == "chk":
                run.bold = True
            elif kind == "txt":
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    figs = annex_figs(figs)
    if figs:
        doc.add_page_break()
        doc.add_heading("Anexo B — Dibujos y resultados graficos", level=1)
        for fp in figs:
            try:
                doc.add_picture(fp, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                pass

    doc.save(path)
    return path


# ===================================================================== PDF
def export_pdf(prj: Project, res: Results, path: str,
               figs: list[str] | None = None, detail: bool = True) -> str:
    """Memoria de calculo en PDF (reportlab, sin depender de Word)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Image as RLImage, PageBreak)

    u = U(prj.metric)
    us = _units(prj)
    ss = getSampleStyleSheet()
    H0 = ParagraphStyle("H0", parent=ss["Title"], fontName="Helvetica-Bold",
                        fontSize=15, spaceAfter=2)
    H1 = ParagraphStyle("H1", parent=ss["Heading2"], fontName="Helvetica-Bold",
                        fontSize=11, spaceBefore=10, spaceAfter=4,
                        textColor=colors.HexColor("#1F3864"))
    BODY = ParagraphStyle("BODY", parent=ss["BodyText"], fontName="Helvetica",
                          fontSize=8, leading=10.5)
    SMALL = ParagraphStyle("SMALL", parent=BODY, fontSize=6.6, leading=8.2,
                           textColor=colors.HexColor("#444444"))
    CEN = ParagraphStyle("CEN", parent=BODY, alignment=TA_CENTER)

    story = []
    try:
        from . import brand
        from reportlab.platypus import Image as RLImage
        _lw = 2.4 * inch
        story.append(RLImage(brand.LOGO(), width=_lw, height=_lw * 246.0 / 1100.0, hAlign="LEFT"))
        story.append(Spacer(1, 4))
    except Exception:
        pass
    story.append(Paragraph("MEMORIA DE CALCULO — PLACA BASE", H0))
    story.append(Paragraph(
        f"<b>{prj.name}</b> &nbsp;|&nbsp; Elemento: {prj.element} &nbsp;|&nbsp; "
        f"Calculo: {prj.author or '-'} &nbsp;|&nbsp; "
        f"Fecha: {prj.date or datetime.date.today().isoformat()}", CEN))
    story.append(Paragraph(
        "Normas: AISC 360-22 · AISC Design Guide 1 (2ª Ed.) · ACI 318-19 Cap. 17"
        f" &nbsp;|&nbsp; Unidades: {us.L}, {us.F}, {us.S}", CEN))
    story.append(Spacer(1, 8))

    def tbl(data, widths, style_extra=None, hdr=True):
        t = Table(data, colWidths=widths, repeatRows=1 if hdr else 0)
        st = [("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
              ("FONTSIZE", (0, 0), (-1, -1), 7),
              ("LEADING", (0, 0), (-1, -1), 8.6),
              ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
              ("TOPPADDING", (0, 0), (-1, -1), 2),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        if hdr:
            st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        t.setStyle(TableStyle(st + (style_extra or [])))
        return t

    # ------------------------------------------------------- 1. entrada
    story.append(Paragraph("1. Datos de entrada", H1))
    import tempfile
    _tmp3d = render_3d_png(prj, str(Path(tempfile.mkdtemp()) / "vista3d.png"))
    if _tmp3d:
        from PIL import Image as _PIL
        _iw, _ih = _PIL.open(_tmp3d).size
        _w = 5.4 * inch
        story.append(RLImage(_tmp3d, width=_w, height=_w * _ih / _iw))
        story.append(Paragraph(f"Vista 3D de la conexion con las cargas de la combinacion {_combo_name(prj)}", CEN))
        story.append(Spacer(1, 4))
    p, b, c, L = prj.plate, prj.bolts, prj.conc, prj.eloads
    geo = (f"Ø{us.q('L', p.Dp)}" if p.shape == "Circular"
           else f"{us.q('L', p.N)} × {us.q('L', p.B)}")
    rows = [["Concepto", "Descripcion"]]
    rows += [
        ["Perfil", f"{prj.section.label} ({prj.section.steel}), rotacion "
                   f"{prj.section.rotation:g}°"],
        ["Placa base", f"{geo} × {us.q('L', p.tp)} — {p.steel} "
                       f"(Fy = {us.q('S', p.mat().Fy)})"],
        ["Mortero de nivelacion", us.q("L", p.grout)],
        ["Anclajes", f"{b.n_total} Ø{b.size} in ({us.q('L', b.geom().db)}) — {b.steel} "
                     f"— {b.atype}, hef = {us.q('L', b.hef)}"],
        ["Agujeros en la placa", f"{us.q('L', b.geom().dh)} — {b.hole_rule}"],
        ["Disposicion", f"{b.pattern}; {b.n_major} en eje mayor, {b.n_minor} en eje "
                        f"menor; ex = {us.q('L', b.ex)}, ey = {us.q('L', b.ey)}"],
        ["Concreto", f"f'c = {us.q('S', c.fc)}; pedestal {us.q('L', c.N2)} × "
                     f"{us.q('L', c.B2)}; ha = {us.q('L', c.ha)}; "
                     f"{'fisurado' if c.cracked else 'no fisurado'}; "
                     f"condicion {'A' if c.cond_A_eff else 'B'}"],
        ["Llave de corte", (f"{us.q('L', prj.lug.W)} × {us.q('L', prj.lug.H)} × "
                            f"{us.q('L', prj.lug.t)} — {prj.lug.direction}"
                            if prj.lug.enabled else "No")],
        ["Rigidizadores", (f"{prj.stiff.count} pletinas {prj.stiff.shape.lower()}, "
                           f"t = {us.q('L', prj.stiff.t)}, L = {us.q('L', prj.stiff.L)}, "
                           f"h = {us.q('L', prj.stiff.h)}; {prj.stiff.position}; "
                           f"{prj.stiff.spacing_mode}"
                           + (f", separacion {us.q('L', prj.stiff.spacing)}"
                              if prj.stiff.spacing_mode.startswith("Separacion") else "")
                           if prj.stiff.enabled else "No")],
        ["Soldadura", f"ala: {prj.welds.flange.wtype} "
                      f"{us.q('L', prj.welds.flange.size)} "
                      f"({prj.welds.flange.electrode}) · alma: {prj.welds.web.wtype} "
                      f"{us.q('L', prj.welds.web.size)} · perimetral: "
                      f"{prj.welds.perimeter.wtype} "
                      f"{us.q('L', prj.welds.perimeter.size)}"],
        ["Cargas (LRFD)", f"Pu = {us.q('F', L.Pu)}; Mux = {us.q('M', L.Mux)}; "
                          f"Muy = {us.q('M', L.Muy)}; Vux = {us.q('F', L.Vux)}; "
                          f"Vuy = {us.q('F', L.Vuy)}"],
    ]
    rows = [[Paragraph(str(a), BODY), Paragraph(str(bq), BODY)] for a, bq in rows]
    story.append(tbl(rows, [1.5 * inch, 5.2 * inch]))

    # --------------------------------------------- 2. equilibrio
    story.append(Paragraph("2. Aplastamiento y equilibrio (DG1 §3.3)", H1))
    br = res.br
    story.append(Paragraph(
        f"A1 = {us.fmt('A', br.A1)} {us.A}; A2 = {us.fmt('A', br.A2)} {us.A}; "
        f"√(A2/A1) = {br.sqrt_ratio:.3f}; φcPp = {us.q('F', br.phiPp)}; "
        f"fp,max = {us.q('S', br.fp_max)}; qmax = {us.q('LF', br.qmax)}.", BODY))
    ee = "∞" if br.e == float("inf") else us.q("L", br.e)
    story.append(Paragraph(
        f"e = Mu/Pu = {ee}; ecrit = {us.q('L', br.ecrit)} → <b>{br.case}</b>. "
        f"Y = {us.q('L', br.Y)}; fp = {us.q('S', br.fp)}; Tu = {us.q('F', br.Tu)} "
        f"repartidos en {br.n_t} pernos con brazo f = {us.q('L', br.f_arm)}.", BODY))
    d = res.tdet
    story.append(Paragraph(
        f"Voladizos: m = {us.q('L', d['m_y'])}, n = {us.q('L', d['m_x'])}, "
        f"λ = {d['lam']:.3f}, λn' = {us.q('L', d['lam_n'])}. "
        f"Espesor requerido t = <b>{us.q('L', res.treq)}</b> "
        f"(propuesto {us.q('L', p.tp)}).", BODY))

    # --------------------------------------------- 3. verificaciones
    story.append(Paragraph("3. Verificaciones", H1))
    data = [["Verificacion", "Demanda", "Capacidad", "Un.", "D/C"]]
    style = []
    for i, ch in enumerate(res.checks, start=1):
        dv, cv, ul = ck_vals(us, ch)
        data.append([Paragraph(ch.title, BODY), f"{dv:,.3f}", f"{cv:,.3f}",
                     ul, "—" if ch.skip else f"{ch.ratio:.3f}"])
        col = colors.HexColor("#EEEEEE") if ch.skip else (
            colors.HexColor("#C6EFCE") if ch.ok else colors.HexColor("#FFC7CE"))
        style.append(("BACKGROUND", (4, i), (4, i), col))
    style += [("ALIGN", (1, 1), (-1, -1), "RIGHT"),
              ("FONTNAME", (4, 1), (4, -1), "Helvetica-Bold")]
    story.append(tbl(data, [3.5 * inch, 0.85 * inch, 0.95 * inch, 0.5 * inch,
                            0.6 * inch], style))

    gov = res.governing
    vcol = "#9C5700" if res.pending else ("#006100" if res.ok else "#9C0006")
    vtxt = "PENDIENTE — realizar analisis 3D" if res.pending else ("CUMPLE" if res.ok else "NO CUMPLE")
    story.append(Spacer(1, 5))
    story.append(Paragraph(
        f'<b>VEREDICTO: <font color="{vcol}">'
        f'{vtxt}</font></b> &nbsp;&nbsp; '
        f"D/C maximo = {res.max_ratio:.3f}"
        + (f" &nbsp;(gobierna: {gov.title})" if gov else ""), BODY))

    if len(getattr(res, "combo_rows", []) or []) > 1:
        story.append(Spacer(1, 4))
        story.append(Paragraph("<b>Combinaciones de carga</b> (el detalle de esta memoria corresponde a la "
                               "que gobierna)", BODY))
        cd = [["Combinacion", "D/C max", "Gobierna", "Estado"]]
        cst = []
        for i, cr in enumerate(res.combo_rows, start=1):
            cd.append([cr["name"], f"{cr['ratio']:.3f}", Paragraph(cr["gov"], BODY),
                       "PENDIENTE" if cr["pending"] else ("CUMPLE" if cr["ok"] else "NO CUMPLE")])
            cst.append(("BACKGROUND", (3, i), (3, i), colors.HexColor(
                "#FFEB9C" if cr["pending"] else ("#C6EFCE" if cr["ok"] else "#FFC7CE"))))
        story.append(tbl(cd, [1.4 * inch, 0.8 * inch, 3.3 * inch, 0.9 * inch], cst))

    if res.warnings:
        story.append(Paragraph("4. Avisos", H1))
        for wmsg in res.warnings:
            col = "9C0006" if wmsg.startswith("**") else "7F6000"
            story.append(Paragraph(f'<font color="#{col}">• {wmsg}</font>', BODY))

    # --------------------------------------------- 6-7. modelo solido 3D
    _fs = fem_section(prj, res, us)
    if _fs:
        story.append(Paragraph("6. Modelo solido 3D — pernos y soldadura", H1))
        story.append(Paragraph(_fs["intro"], BODY))
        story.append(Paragraph(_fs["msg"], BODY))
        wr = _fs["weld_rows"]
        st = []
        for i, r in enumerate(wr[1:], start=1):
            for j in (5, 6):
                try:
                    ok = float(r[j]) <= (float(getattr(prj.fea, "weld_peak_factor", 1.5)) if j == 5 else 1.0)
                except ValueError:
                    ok = False
                st.append(("BACKGROUND", (j, i), (j, i),
                           colors.HexColor("#C6EFCE" if ok else "#FFC7CE")))
        story.append(Spacer(1, 4))
        story.append(Paragraph("<b>Soldadura perfil-placa</b>", BODY))
        story.append(tbl([[Paragraph(str(c), BODY) for c in r] for r in wr],
                         [0.8 * inch, 1.9 * inch] + [0.8 * inch] * 5, st))
        if _fs["b_hdr"]:
            story.append(Spacer(1, 4))
            story.append(Paragraph("<b>Traccion por perno</b>", BODY))
            bst = []
            for k_, r in enumerate(_fs["b_rows"], start=1):
                try:
                    okb = float(r[4]) <= 1.0
                except ValueError:
                    continue
                bst.append(("BACKGROUND", (4, k_), (4, k_),
                            colors.HexColor("#C6EFCE" if okb else "#FFC7CE")))
            story.append(tbl([_fs["b_hdr"]] + _fs["b_rows"], [0.6 * inch] + [1.0 * inch] * 4,
                             [("ALIGN", (1, 1), (-1, -1), "RIGHT")] + bst))
        story.append(PageBreak())
        story.append(Paragraph("7. Placa base aislada — resultados del modelo 3D en planta", H1))
        for ln in _fs["lines"]:
            story.append(Paragraph(f"<b>{ln}</b>" if ln.startswith("ESFUERZO MAXIMO") else ln, BODY))
        from PIL import Image as _PIL2
        for fp, cap in _fs["figs"] + _fs["persp"]:
            try:
                iw, ih = _PIL2.open(fp).size
                w = 4.9 * inch
                story.append(Spacer(1, 6))
                story.append(RLImage(fp, width=w, height=w * ih / iw))
                story.append(Paragraph(cap, CEN))
            except Exception:
                pass

    # ------------------------------------- desarrollo de las ecuaciones
    if detail and res.rec is not None:
        story.append(PageBreak())
        story.append(Paragraph("Anexo A — Desarrollo de las ecuaciones", H1))
        story.append(Paragraph(
            "Cada linea muestra el simbolo, la formula, la sustitucion numerica y "
            "el resultado, en las unidades de trabajo del proyecto.", SMALL))
        EQ = ParagraphStyle("EQ", parent=BODY, fontName="Helvetica",
                            fontSize=7.4, leading=9.6, leftIndent=8)
        SEC = ParagraphStyle("SEC", parent=BODY, fontName="Helvetica-Bold",
                             fontSize=8.4, leading=11, spaceBefore=7, spaceAfter=2,
                             textColor=colors.white,
                             backColor=colors.HexColor("#1F3864"),
                             leftIndent=2, rightIndent=2, borderPadding=3)
        NT = ParagraphStyle("NT", parent=EQ, textColor=colors.HexColor("#555555"),
                            fontName="Helvetica-Oblique", leftIndent=16)
        CH = ParagraphStyle("CH", parent=EQ, fontName="Helvetica-Bold",
                            textColor=colors.HexColor("#1F3864"))
        from xml.sax.saxutils import escape as _esc
        for kind, txt in res.rec.to_lines():
            st = {"sec": SEC, "txt": NT, "chk": CH}.get(kind, EQ)
            story.append(Paragraph(_esc(txt), st))

    # --------------------------------------------- anexo grafico
    figs = annex_figs(figs)
    if figs:
        story.append(PageBreak())
        story.append(Paragraph("Anexo — Dibujos y resultados graficos", H1))
        for fp in figs:
            try:
                from PIL import Image as PILImage
                iw, ih = PILImage.open(fp).size
                w = 6.1 * inch
                story.append(RLImage(fp, width=w, height=w * ih / iw))
                story.append(Spacer(1, 6))
            except Exception:
                pass

    doc = SimpleDocTemplate(path, pagesize=letter,
                            leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.55 * inch, bottomMargin=0.55 * inch,
                            title=f"Memoria placa base {prj.element}",
                            author=prj.author or "PlacaBasePro")

    def _footer(canvas, docu):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(0.6 * inch, 0.32 * inch,
                          f"{prj.name} — {prj.element} — PlacaBasePro {__import__('placabase').__version__}")
        canvas.drawRightString(letter[0] - 0.6 * inch, 0.32 * inch,
                               f"Pagina {docu.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return path
