# -*- coding: utf-8 -*-
"""Memoria de calculo (PDF y Word) de las tipologias de conexion distintas de la placa base.

`report.py` delega aqui cuando `prj.ctype` no es la placa base.  Misma estructura que la memoria de la placa base:
datos, combinaciones, verificaciones, veredicto, avisos, anexo de ecuaciones y dibujos.
"""
from __future__ import annotations
import datetime

from ..units import UnitSet

def _us(prj) -> UnitSet:
    return UnitSet(prj.u_len, prj.u_force, prj.u_stress, prj.u_moment)


def _mod(prj):
    from . import module_for
    return module_for(prj.ctype)


def input_rows(prj, us: UnitSet) -> list:
    """Filas (concepto, descripcion) de los datos de entrada."""
    return _mod(prj).input_rows(prj, us)


def _checks_rows(us, res):
    from ..report import ck_vals
    out = []
    for ch in res.checks:
        dv, cv, ul = ck_vals(us, ch)
        out.append((ch, dv, cv, ul))
    return out


def _verdict(res):
    if res.pending:
        return "PENDIENTE", "#9C5700"
    return ("CUMPLE", "#006100") if res.ok else ("NO CUMPLE", "#9C0006")


# ======================================================================= PDF
def export_pdf(prj, res, path: str, figs: list | None = None, detail: bool = True) -> str:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage,
                                    PageBreak)
    from xml.sax.saxutils import escape as esc

    us = _us(prj)
    ss = getSampleStyleSheet()
    H0 = ParagraphStyle("H0", parent=ss["Title"], fontName="Helvetica-Bold", fontSize=15, spaceAfter=2)
    H1 = ParagraphStyle("H1", parent=ss["Heading2"], fontName="Helvetica-Bold", fontSize=11, spaceBefore=10,
                        spaceAfter=4, textColor=colors.HexColor("#1F3864"))
    BODY = ParagraphStyle("BODY", parent=ss["BodyText"], fontName="Helvetica", fontSize=8, leading=10.5)
    SMALL = ParagraphStyle("SMALL", parent=BODY, fontSize=6.6, leading=8.2, textColor=colors.HexColor("#444444"))
    CEN = ParagraphStyle("CEN", parent=BODY, alignment=TA_CENTER)

    def tbl(data, widths, extra=None, hdr=True):
        t = Table(data, colWidths=widths, repeatRows=1 if hdr else 0)
        st = [("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 7),
              ("LEADING", (0, 0), (-1, -1), 8.6), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 2),
              ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]
        if hdr:
            st += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E75B6")),
                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
        t.setStyle(TableStyle(st + (extra or [])))
        return t

    story = []
    try:
        from .. import brand
        lw = 2.4 * inch
        story.append(RLImage(brand.LOGO(), width=lw, height=lw * 246.0 / 1100.0, hAlign="LEFT"))
        story.append(Spacer(1, 4))
    except Exception:
        pass
    story.append(Paragraph(f"MEMORIA DE CALCULO — {_mod(prj).TITLE}", H0))
    story.append(Paragraph(f"<b>{esc(prj.name)}</b> &nbsp;|&nbsp; Elemento: {esc(prj.element)} &nbsp;|&nbsp; "
                           f"Calculo: {esc(prj.author or '-')} &nbsp;|&nbsp; "
                           f"Fecha: {prj.date or datetime.date.today().isoformat()}", CEN))
    story.append(Paragraph(f"Normas: {_mod(prj).NORMS} &nbsp;|&nbsp; Unidades: {us.L}, {us.F}, {us.S}", CEN))
    story.append(Spacer(1, 8))

    story.append(Paragraph("1. Datos de entrada", H1))
    rows = [["Concepto", "Descripcion"]] + [[Paragraph(esc(a), BODY), Paragraph(esc(b), BODY)]
                                           for a, b in input_rows(prj, us)]
    story.append(tbl(rows, [1.5 * inch, 5.2 * inch]))

    story.append(Paragraph("2. Verificaciones", H1))
    data = [["Verificacion", "Demanda", "Capacidad", "Un.", "D/C"]]
    style = []
    for i, (ch, dv, cv, ul) in enumerate(_checks_rows(us, res), start=1):
        data.append([Paragraph(esc(ch.title), BODY), f"{dv:,.3f}", f"{cv:,.3f}", ul,
                     "—" if ch.skip else f"{ch.ratio:.3f}"])
        style.append(("BACKGROUND", (4, i), (4, i), colors.HexColor(
            "#EEEEEE" if ch.skip else ("#C6EFCE" if ch.ok else "#FFC7CE"))))
    style += [("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("FONTNAME", (4, 1), (4, -1), "Helvetica-Bold")]
    story.append(tbl(data, [3.5 * inch, 0.85 * inch, 0.95 * inch, 0.5 * inch, 0.6 * inch], style))
    gov = res.governing
    vtxt, vcol = _verdict(res)
    story.append(Spacer(1, 5))
    story.append(Paragraph(f'<b>VEREDICTO: <font color="{vcol}">{vtxt}</font></b> &nbsp;&nbsp; '
                           f"D/C maximo = {res.max_ratio:.3f}" + (f" &nbsp;(gobierna: {esc(gov.title)})" if gov else ""),
                           BODY))
    if len(res.combo_rows or []) > 1:
        story.append(Spacer(1, 4))
        story.append(Paragraph("<b>Combinaciones de carga</b> (el detalle corresponde a la que gobierna)", BODY))
        cd = [["Combinacion", "D/C max", "Gobierna", "Estado"]]
        cst = []
        for i, cr in enumerate(res.combo_rows, start=1):
            cd.append([cr["name"], f"{cr['ratio']:.3f}", Paragraph(esc(cr["gov"]), BODY),
                       "CUMPLE" if cr["ok"] else "NO CUMPLE"])
            cst.append(("BACKGROUND", (3, i), (3, i), colors.HexColor("#C6EFCE" if cr["ok"] else "#FFC7CE")))
        story.append(tbl(cd, [1.4 * inch, 0.8 * inch, 3.3 * inch, 0.9 * inch], cst))

    if res.warnings:
        story.append(Paragraph("3. Avisos", H1))
        for w in res.warnings:
            col = "9C0006" if w.startswith("**") else "7F6000"
            story.append(Paragraph(f'<font color="#{col}">• {esc(w)}</font>', BODY))

    if detail and res.rec is not None:
        story.append(PageBreak())
        story.append(Paragraph("Anexo A — Desarrollo de las ecuaciones", H1))
        story.append(Paragraph("Cada linea muestra el simbolo, la formula, la sustitucion numerica y el resultado, "
                               "en las unidades de trabajo del proyecto.", SMALL))
        EQ = ParagraphStyle("EQ", parent=BODY, fontSize=7.4, leading=9.6, leftIndent=8)
        SEC = ParagraphStyle("SEC", parent=BODY, fontName="Helvetica-Bold", fontSize=8.4, leading=11, spaceBefore=7,
                             spaceAfter=2, textColor=colors.white, backColor=colors.HexColor("#1F3864"),
                             leftIndent=2, rightIndent=2, borderPadding=3)
        NT = ParagraphStyle("NT", parent=EQ, textColor=colors.HexColor("#555555"), fontName="Helvetica-Oblique",
                            leftIndent=16)
        CH = ParagraphStyle("CH", parent=EQ, fontName="Helvetica-Bold", textColor=colors.HexColor("#1F3864"))
        for kind, txt in res.rec.to_lines():
            story.append(Paragraph(esc(txt), {"sec": SEC, "txt": NT, "chk": CH}.get(kind, EQ)))

    if figs:
        story.append(PageBreak())
        story.append(Paragraph("Anexo B — Dibujos", H1))
        from PIL import Image as PILImage
        for fp in figs:
            try:
                iw, ih = PILImage.open(fp).size
                w = 6.1 * inch
                story.append(RLImage(fp, width=w, height=w * ih / iw))
                story.append(Spacer(1, 6))
            except Exception:
                pass

    doc = SimpleDocTemplate(path, pagesize=letter, leftMargin=0.6 * inch, rightMargin=0.6 * inch,
                            topMargin=0.55 * inch, bottomMargin=0.55 * inch,
                            title=f"Memoria {_mod(prj).TITLE.lower()} {prj.element}",
                            author=prj.author or "SteelConPro")

    def _footer(canvas, docu):
        canvas.saveState()
        canvas.setFont("Helvetica", 6.5)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawString(0.6 * inch, 0.32 * inch,
                          f"{prj.name} — {prj.element} — SteelConPro {__import__('steelconpro').__version__}")
        canvas.drawRightString(letter[0] - 0.6 * inch, 0.32 * inch, f"Pagina {docu.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return path


# ====================================================================== Word
def export_docx(prj, res, path: str, figs: list | None = None, detail: bool = True) -> str:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    us = _us(prj)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(9)
    try:
        from .. import brand
        doc.add_picture(brand.LOGO(), width=Inches(2.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    except Exception:
        pass
    doc.add_heading(f"MEMORIA DE CALCULO — {_mod(prj).TITLE}", level=0)
    p0 = doc.add_paragraph()
    p0.add_run(f"{prj.name}\n").bold = True
    p0.add_run(f"Elemento: {prj.element}     Calculo: {prj.author}     "
               f"Fecha: {prj.date or datetime.date.today().isoformat()}\n")
    p0.add_run(f"Normas: {_mod(prj).NORMS}.")

    doc.add_heading("1. Datos de entrada", level=1)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Grid Accent 1"
    for k, v in input_rows(prj, us):
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = v

    doc.add_heading("2. Verificaciones", level=1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ("Verificacion", "Demanda", "Capacidad", "Un.", "D/C")):
        c.text = h
    for ch, dv, cv, ul in _checks_rows(us, res):
        cells = t.add_row().cells
        cells[0].text = ch.title
        cells[1].text = f"{dv:,.3f}"
        cells[2].text = f"{cv:,.3f}"
        cells[3].text = ul
        cells[4].text = "—" if ch.skip else f"{ch.ratio:.3f}"
    vtxt, vcol = _verdict(res)
    gov = res.governing
    p = doc.add_paragraph()
    p.add_run("VEREDICTO: ").bold = True
    r = p.add_run(vtxt)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(vcol.lstrip("#"))
    p.add_run(f"     D/C maximo = {res.max_ratio:.3f}" + (f"  (gobierna: {gov.title})" if gov else ""))
    if len(res.combo_rows or []) > 1:
        doc.add_paragraph("Combinaciones de carga (el detalle corresponde a la que gobierna):")
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ("Combinacion", "D/C max", "Gobierna", "Estado")):
            c.text = h
        for cr in res.combo_rows:
            cells = t.add_row().cells
            cells[0].text, cells[1].text = cr["name"], f"{cr['ratio']:.3f}"
            cells[2].text, cells[3].text = cr["gov"], "CUMPLE" if cr["ok"] else "NO CUMPLE"
    if res.warnings:
        doc.add_heading("3. Avisos", level=1)
        for w in res.warnings:
            doc.add_paragraph(w, style="List Bullet")
    if detail and res.rec is not None:
        doc.add_heading("Anexo A — Desarrollo de las ecuaciones", level=1)
        for kind, txt in res.rec.to_lines():
            if kind == "sec":
                doc.add_heading(txt, level=2)
            else:
                para = doc.add_paragraph(txt)
                para.paragraph_format.space_after = Pt(1)
                if kind == "chk":
                    for run in para.runs:
                        run.bold = True
                elif kind == "txt":
                    for run in para.runs:
                        run.italic = True
    if figs:
        doc.add_heading("Anexo B — Dibujos", level=1)
        for fp in figs:
            try:
                doc.add_picture(fp, width=Inches(6.0))
            except Exception:
                pass
    doc.save(path)
    return path
